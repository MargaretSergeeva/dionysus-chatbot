#!/usr/bin/env python3
"""Reviewer Excel: build the blank template, fill it from a run, or import a reviewed file back to CSV.

  template                       -> evaluation/templates/review_template.xlsx (all question IDs, no answers)
  fill --run <run.csv> [--split] -> evaluation/review/<run>.xlsx (answers filled; --split = one file per reviewer)
  import <reviewed.xlsx> --out <csv>  -> validated CSV (dropdown values checked) to mirror into GitHub

Columns: locked = question_id, category, question, answer, reviewer; editable = score, ideal_answer, fix_type,
check_length, check_link_at_end, check_general_vs_specific, comment.
SCORE IS INVERTED: 1 = super ... 5 = bad (stated in the header cell comment and on the Legende sheet).
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.worksheet.datavalidation import DataValidation

from common import CHECKS, FIX_TYPES, REVIEW, REVIEWERS, TEMPLATE, load_questions, read_table, write_csv

COLS = [
    # name, width, locked, header
    ("question_id", 11, True, "question_id"),
    ("category", 22, True, "category"),
    ("question", 48, True, "question"),
    ("answer", 60, True, "answer"),
    ("reviewer", 12, True, "reviewer"),
    ("score", 14, False, "score (1 = super, 5 = bad)"),
    ("ideal_answer", 55, False, "ideal_answer"),
    ("fix_type", 12, False, "fix_type"),
    ("check_length", 13, False, "check_length"),
    ("check_link_at_end", 17, False, "check_link_at_end"),
    ("check_general_vs_specific", 22, False, "check_general_vs_specific"),
    ("comment", 40, False, "comment"),
]
NAMES = [c[0] for c in COLS]
HELP = {
    "score": "INVERTED scale: 1 = super, 2 = good, 3 = ok, 4 = weak, 5 = bad.",
    "fix_type": "Where the fix goes: prompt (Oksana) / Vova (Gastbot/dev) / data (missing or wrong source) / links.",
    "check_length": "Rozaliia rule: answer length is limited. ok / fail / n/a.",
    "check_link_at_end": "Rozaliia rule: a general link on the topic at the end. ok / fail / n/a.",
    "check_general_vs_specific": "Rozaliia rule: general question -> general answer, specific question -> specific answer. ok / fail / n/a.",
}


def build(rows: list[dict], answers: dict[str, str] | None, path: Path, reviewer: str | None = None) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Review"
    head = PatternFill("solid", fgColor="DDDDDD")
    edit = PatternFill("solid", fgColor="FFF9DB")
    for i, (name, width, _, header) in enumerate(COLS, 1):
        c = ws.cell(row=1, column=i, value=header)
        c.font = Font(bold=True)
        c.fill = head
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[c.column_letter].width = width
        if name in HELP:
            c.comment = Comment(HELP[name], "Gastbot review")
    ws.freeze_panes = "D2"
    for r, q in enumerate(rows, 2):
        vals = {"question_id": q["id"], "category": q["category"], "question": q["question"],
                "answer": (answers or {}).get(q["id"], ""), "reviewer": q.get("reviewer", "")}
        for i, (name, _, locked, _) in enumerate(COLS, 1):
            c = ws.cell(row=r, column=i, value=vals.get(name, ""))
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.protection = Protection(locked=locked)
            if not locked:
                c.fill = edit
    last = len(rows) + 1

    def dv(kind_formula, name, error):
        col = chr(ord("A") + NAMES.index(name))
        v = DataValidation(type="list", formula1=kind_formula, allow_blank=True, showErrorMessage=True, error=error)
        v.add(f"{col}2:{col}{last}")
        ws.add_data_validation(v)

    dv('"1,2,3,4,5"', "score", "Score 1-5 (1 = super, 5 = bad)")
    dv('"' + ",".join(FIX_TYPES) + '"', "fix_type", "prompt / Vova / data / links")
    for chk in ("check_length", "check_link_at_end", "check_general_vs_specific"):
        dv('"' + ",".join(CHECKS) + '"', chk, "ok / fail / n/a")
    ws.protection.sheet = True                  # no password: protects against accidents, not against people
    ws.protection.formatColumns = False
    ws.protection.formatRows = False
    ws.auto_filter.ref = f"A1:{chr(ord('A') + len(COLS) - 1)}{last}"

    lg = wb.create_sheet("Legende")
    lines = ["How to review", "",
             "SCORE IS INVERTED: 1 = super ... 5 = bad.", "",
             "Fill only the yellow columns. Review only rows where reviewer = your name.",
             "ideal_answer: what the bot should have said (short).",
             "fix_type: prompt = Oksana's prompt | Vova = Gastbot/dev | data = missing/wrong source | links = link entries.", "",
             "Prompt rules (from Rozaliia) checked per row:",
             " - check_length: answer length is limited",
             " - check_link_at_end: a general link on the topic at the end",
             " - check_general_vs_specific: general question -> general answer; specific -> specific", "",
             "Do not change question_id, answer or reviewer (locked). Question IDs are the same in every run."]
    if reviewer:
        lines.insert(2, f"This file: rows for {reviewer}.")
    for i, t in enumerate(lines, 1):
        lg.cell(row=i, column=1, value=t)
    lg["A1"].font = Font(bold=True)
    lg.column_dimensions["A"].width = 110
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def cmd_template(_):
    build(load_questions(), None, TEMPLATE)
    print(TEMPLATE)


def cmd_fill(a):
    qs = {q["id"]: q for q in load_questions()}
    run = read_table(a.run)
    answers = {r["question_id"]: r["answer"] for r in run}
    rows = [qs[i] for i in sorted(qs) if i in answers]
    stem = REVIEW / a.run.name.removesuffix('.csv')
    if a.split:
        for who in REVIEWERS:
            p = REVIEW / f"{stem.name}__{who}.xlsx"
            build([q for q in rows if q["reviewer"] == who], answers, p, reviewer=who)
            print(p)
    else:
        p = REVIEW / f"{stem.name}.xlsx"
        build(rows, answers, p)
        print(p)


def cmd_import(a):
    rows = [{next((n for n in NAMES if k == n or k.startswith(n + " (")), k): v for k, v in r.items()} for r in read_table(a.file)]
    known = {q["id"] for q in load_questions()}
    errs = []
    for r in rows:
        i = r.get("question_id", "")
        if i not in known:
            errs.append(f"unknown question_id {i!r}")
        s = r.get("score", "").strip()
        if s and s.split(".")[0] not in "12345":
            errs.append(f"{i}: score {s!r} not 1-5")
        if r.get("fix_type") and r["fix_type"] not in FIX_TYPES:
            errs.append(f"{i}: fix_type {r['fix_type']!r}")
        for c in ("check_length", "check_link_at_end", "check_general_vs_specific"):
            if r.get(c) and r[c] not in CHECKS:
                errs.append(f"{i}: {c} {r[c]!r}")
        if r.get("score", "").strip() in ("4", "5") and not r.get("fix_type"):
            errs.append(f"{i}: score {r['score']} but no fix_type")
    if errs:
        sys.exit("\n".join(errs))
    write_csv(a.out, rows, NAMES)
    done = sum(1 for r in rows if r.get("score", "").strip())
    print(f"{a.out}: {len(rows)} rows, {done} scored")


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    s = p.add_subparsers(dest="cmd", required=True)
    s.add_parser("template").set_defaults(fn=cmd_template)
    f = s.add_parser("fill")
    f.add_argument("--run", type=Path, required=True)
    f.add_argument("--split", action="store_true")
    f.set_defaults(fn=cmd_fill)
    i = s.add_parser("import")
    i.add_argument("file", type=Path)
    i.add_argument("--out", type=Path, required=True)
    i.set_defaults(fn=cmd_import)
    a = p.parse_args()
    a.fn(a)


if __name__ == "__main__":
    main()
